"""test_import_export.py — HTTP API tests for /master-import/export{,.xlsx}.

IMPORTANT: does NOT import main.app (which would trigger start_scheduler() against
prod Atlas). Builds a minimal FastAPI app with only the dynamic-import router,
mirroring tests/test_import_endpoints.py.
"""
import io
import os
import uuid

import pytest
import pytest_asyncio
from motor.motor_asyncio import AsyncIOMotorClient
from httpx import AsyncClient, ASGITransport
from fastapi import FastAPI
from openpyxl import load_workbook

# Verify we are not running against prod
_DB_NAME = os.getenv("DB_NAME", "smartshape_test")
assert _DB_NAME.endswith("_test") or _DB_NAME == "mtt_ci", (
    f"refusing non-test DB: {_DB_NAME}"
)

import routes.dynamic_import_routes as dir_mod
import auth_utils
import field_registry as fr

# Build minimal app — no scheduler, no startup hooks
_app = FastAPI()
_app.include_router(dir_mod.router, prefix="/api")
_app.dependency_overrides[auth_utils.get_current_user] = lambda: {
    "email": "admin@t",
    "team": "admin",
    "role": "admin",
    "assigned_modules": ["settings"],
}

_MONGO_URL = os.getenv("MONGO_URL", "mongodb://localhost:27017")


@pytest_asyncio.fixture
async def test_db():
    """Fresh Motor db handle created INSIDE the current test's event loop.

    Also patches dir_mod.db so the route handlers use this same connection,
    and seeds field definitions. Tears down all test collections after each test.
    """
    motor_client = AsyncIOMotorClient(_MONGO_URL)
    d = motor_client[_DB_NAME]

    await fr.seed_field_definitions(d)

    original_db = dir_mod.db
    dir_mod.db = d

    yield d

    dir_mod.db = original_db
    await d.field_definitions.delete_many({})
    await d.app_meta.delete_many({"_id": "field_definitions_seeded"})
    await d.import_logs.delete_many({})
    await d.schools.delete_many({})
    await d.contacts.delete_many({})
    await d.audit_backup.delete_many({})
    motor_client.close()


@pytest_asyncio.fixture
async def client(test_db):
    """Async HTTP client wrapping the minimal app — shares the test event loop."""
    transport = ASGITransport(app=_app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        yield ac


# ---------------------------------------------------------------------------
# /master-import/export tests
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_export_includes_school_contact_and_custom(client, test_db):
    # Custom school field
    r_field = await client.post("/api/fields", json={
        "label": "Transport Fee",
        "entity": "school",
        "type": "number",
    })
    assert r_field.status_code == 200

    sid = f"sch_{uuid.uuid4().hex[:8]}"
    school = {
        "school_id": sid,
        "school_name": "Sunrise Academy",
        "city": "Pune",
        "is_deleted": False,
        "custom_fields": {"transport_fee": "1500"},
    }
    await test_db.schools.insert_one(school)

    contact = {
        "contact_id": f"con_{uuid.uuid4().hex[:8]}",
        "school_id": sid,
        "name": "Mr Sharma",
        "phone": "9998887770",
        "is_deleted": False,
    }
    await test_db.contacts.insert_one(contact)

    r = await client.get("/api/master-import/export")
    assert r.status_code == 200
    body = r.json()

    assert "School ID" in body["headers"]
    assert "Transport Fee" in body["headers"]

    row = next(r for r in body["rows"] if r["School ID"] == sid)
    assert row["School/Institute Name"] == "Sunrise Academy"
    assert row["Name"] == "Mr Sharma"
    assert row["Transport Fee"] == "1500"


@pytest.mark.asyncio
async def test_export_round_trips_through_import(client, test_db):
    sid = f"sch_{uuid.uuid4().hex[:8]}"
    school = {
        "school_id": sid,
        "school_name": "Greenfield High",
        "city": "Goa",
        "is_deleted": False,
        "custom_fields": {},
    }
    await test_db.schools.insert_one(school)

    r = await client.get("/api/master-import/export")
    assert r.status_code == 200
    body = r.json()
    row = next(r for r in body["rows"] if r["School ID"] == sid)

    # Re-execute directly with the exported row's School ID + name — must
    # resolve to an UPDATE (no duplicate school created).
    keyed = [{
        "school_id": row["School ID"],
        "school_name": row["School/Institute Name"],
        "city": row["City"],
    }]
    r2 = await client.post("/api/master-import/execute", json={
        "rows_keyed": keyed,
        "mapping": [],
        "create_leads": False,
    })
    assert r2.status_code == 200
    assert r2.json()["counts"]["update"] == 1

    count = await test_db.schools.count_documents({"school_name": "Greenfield High"})
    assert count == 1, f"Expected 1 school, found {count}"


@pytest.mark.asyncio
async def test_export_xlsx_returns_spreadsheet(client, test_db):
    sid = f"sch_{uuid.uuid4().hex[:8]}"
    await test_db.schools.insert_one({
        "school_id": sid,
        "school_name": "XLSX School",
        "city": "Delhi",
        "is_deleted": False,
    })

    r = await client.get("/api/master-import/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    first_row = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    r_json = await client.get("/api/master-import/export")
    assert first_row == r_json.json()["headers"]
